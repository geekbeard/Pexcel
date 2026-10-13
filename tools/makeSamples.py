"""Creates sample files in samples/ for manual testing (needs openpyxl from requirements-dev.txt)."""

import datetime
import os
import random

from openpyxl import Workbook

SAMPLE_FOLDER = os.path.join(os.path.dirname(__file__), "..", "samples")


def main():
    os.makedirs(SAMPLE_FOLDER, exist_ok=True)
    random.seed(7)
    workbook = Workbook()

    people = workbook.active
    people.title = "People"
    people.append(["Name", "City", "Age", "Joined", "Salary", "Active", "Notes"])
    cities = ["Amsterdam", "Berlin", "Paris", "Madrid", "Rome", "", "Lisbon"]
    for index in range(1, 501):
        joined = datetime.date(2015, 1, 1) + datetime.timedelta(days=random.randint(0, 3500))
        notes = "" if index % 4 else "Long note text that goes on for a while to test tooltips and eliding " * 2
        people.append([
            "Person {}".format(index), random.choice(cities), random.randint(18, 70),
            joined, round(random.uniform(30000, 120000), 2), index % 3 != 0, notes,
        ])

    report = workbook.create_sheet("Report")
    report["A1"] = "Quarterly report"
    report.merge_cells("A1:D1")
    report.append([])
    report.append(["Quarter", "Revenue", "Cost", "Profit"])
    for quarter in range(1, 5):
        revenue = random.randint(100000, 200000)
        cost = random.randint(50000, 90000)
        report.append(["Q{}".format(quarter), revenue, cost, revenue - cost])

    big = workbook.create_sheet("Big")
    big.append(["Id", "Group", "Value", "Timestamp", "Label"])
    groups = ["Alpha", "Beta", "Gamma", "Delta"]
    start = datetime.datetime(2024, 1, 1, 8, 0)
    for index in range(100000):
        big.append([
            index, random.choice(groups), random.random() * 1000,
            start + datetime.timedelta(minutes=index), "Item {}".format(random.randint(1, 5000)),
        ])

    workbook.create_sheet("Empty")
    hidden = workbook.create_sheet("Secret")
    hidden.append(["Hidden", "Data"])
    hidden.append([1, 2])
    hidden.sheet_state = "hidden"

    path = os.path.abspath(os.path.join(SAMPLE_FOLDER, "sample.xlsx"))
    workbook.save(path)
    print("Wrote " + path)

    csvPath = os.path.abspath(os.path.join(SAMPLE_FOLDER, "sample.csv"))
    with open(csvPath, "w", encoding="utf-8") as file:
        file.write("Code;Product;Price;Stock\n")
        file.write("007;Widget;12,50;3\n")
        file.write("010;Gadget;9.99;12\n")
        file.write("100;Thing;100;\n")
    print("Wrote " + csvPath)


if __name__ == "__main__":
    main()
